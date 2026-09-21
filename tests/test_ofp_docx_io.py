"""Проверка чтения "сырых" отчётов лаборатории по ОФП в формате Word (.docx)."""

import docx

from jfunction.ofp_docx_io import load_ofp_data_from_docx


def _add_meta_table(doc, well, model_num, porosity, perm, swir, sor):
    rows = [
        ("Данные по керну (модели образцов)", "Данные по керну (модели образцов)"),
        ("№ модели образцов", f"Модель №{model_num}"),
        ("Месторождение", "Жанаталап"),
        ("Скважина", well),
        ("Пористость, %", porosity),
        ("Проницаемость по газу, мД", perm),
        ("Полученные результаты", "Полученные результаты"),
        ("Остаточная водонасыщенность (Swi), доли ед.", swir),
        ("Остаточная нефтенасыщенность (Sow), доли ед.", sor),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.rows[0].cells[0].text = "Наименование"
    table.rows[0].cells[1].text = "Значение"
    for i, (label, value) in enumerate(rows, start=1):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value


def _add_curve_table(doc, points):
    table = doc.add_table(rows=1 + len(points), cols=5)
    header = table.rows[0].cells
    header[0].text = "Qнефти, мл/мин"
    header[1].text = "Qводы, мл/мин"
    header[2].text = "Sw"
    header[3].text = "krw"
    header[4].text = "krow"
    for i, (sw, krw, krow) in enumerate(points, start=1):
        cells = table.rows[i].cells
        cells[2].text = sw
        cells[3].text = krw
        cells[4].text = krow


def test_reads_single_model_block(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Скважина №300")
    _add_meta_table(doc, well="300", model_num="1", porosity="34,82", perm="112,50", swir="0,262", sor="0,224")
    doc.add_paragraph("Результаты относительной проницаемости для системы вода-нефть образца модели №1")
    _add_curve_table(
        doc,
        [
            ("0,26", "0,000", "1"),
            ("0,35", "0,013", "0,615"),
            ("0,78", "0,308", "0,000"),
        ],
    )
    path = tmp_path / "ofp.docx"
    doc.save(path)

    df = load_ofp_data_from_docx(path)

    assert len(df) == 3
    assert df["model"].unique().tolist() == ["300-1"]
    assert df["well"].iloc[0] == "300"
    assert df["Swir"].iloc[0] == 0.262
    assert df["Sor"].iloc[0] == 0.224
    assert df["Swmax"].iloc[0] == 0.78
    assert df["krwmax"].iloc[0] == 0.308
    assert df["krow_swc"].iloc[0] == 1.0
    assert df["porosity_pct"].iloc[0] == 34.82
    assert df["perm_mD"].iloc[0] == 112.50


def test_two_curve_tables_share_one_metadata_block(tmp_path):
    """Один керн - две модели, измеренные вместе (как для скв. 305 в реальном отчёте)."""
    doc = docx.Document()
    doc.add_paragraph("Скважина №305")
    _add_meta_table(doc, well="305", model_num="2", porosity="30,0", perm="50,0", swir="0,25", sor="0,20")
    doc.add_paragraph("Результаты ... образца (сэмпл J01H)")
    _add_curve_table(doc, [("0,25", "0,000", "1"), ("0,80", "0,300", "0,000")])
    doc.add_paragraph("Результаты ... образца (сэмпл J03H)")
    _add_curve_table(doc, [("0,25", "0,000", "1"), ("0,75", "0,250", "0,000")])
    path = tmp_path / "ofp_two.docx"
    doc.save(path)

    df = load_ofp_data_from_docx(path)

    assert sorted(df["model"].unique()) == ["305-2", "305-2b"]
    assert len(df) == 4


def _add_flat_table(doc, rows):
    """Строит "плоскую" таблицу формата 2: Горизонт|Скважина|Идентификатор образца|Sw|krw|krow."""
    table = doc.add_table(rows=1 + len(rows), cols=6)
    header = table.rows[0].cells
    header[0].text = "Горизонт"
    header[1].text = "Скважина"
    header[2].text = "Идентификатор образца"
    header[3].text = "Sw"
    header[4].text = "krw"
    header[5].text = "krow"
    for i, (horizon, well, sample, sw, krw, krow) in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text = horizon
        cells[1].text = well
        cells[2].text = sample
        cells[3].text = sw
        cells[4].text = krw
        cells[5].text = krow


def test_reads_flat_table_format_grouped_by_sample(tmp_path):
    """Формат 2: одна таблица Горизонт/Скважина/Идентификатор образца/Sw/krw/krow, без метаданных."""
    doc = docx.Document()
    _add_flat_table(
        doc,
        [
            ("Валанжин", "700", "Модель 1", "0.412", "0.000", "1.000"),
            ("Валанжин", "700", "Модель 1", "0.510", "0.011", "0.554"),
            ("Валанжин", "700", "Модель 1", "0.709", "0.155", "0.000"),
            ("Валанжин", "700", "Модель 1", "1.000", "", ""),
            ("Юра", "НСВ-1", "Модель 5", "0.300", "0.000", "1.000"),
            ("Юра", "НСВ-1", "Модель 5", "0.600", "0.200", "0.000"),
        ],
    )
    path = tmp_path / "ofp_flat.docx"
    doc.save(path)

    df = load_ofp_data_from_docx(path)

    assert set(df["horizon"]) == {"Валанжин", "Юра"}
    valanzhin = df[df["horizon"] == "Валанжин"]
    assert len(valanzhin) == 3  # терминальная строка Sw=1.0 без krw/krow отброшена
    assert valanzhin["model"].unique().tolist() == ["700-1"]
    assert valanzhin["Swir"].iloc[0] == 0.412
    assert valanzhin["Swmax"].iloc[0] == 0.709
    assert valanzhin["krwmax"].iloc[0] == 0.155
    assert valanzhin["krow_swc"].iloc[0] == 1.0
    assert valanzhin["Sor"].iloc[0] == 1 - 0.709

    yura = df[df["horizon"] == "Юра"]
    assert len(yura) == 2
    assert yura["model"].unique().tolist() == ["НСВ-1-5"]


def test_flat_table_skips_samples_with_too_few_points(tmp_path):
    doc = docx.Document()
    _add_flat_table(doc, [("Валанжин", "700", "Модель 1", "0.412", "0.000", "1.000")])
    path = tmp_path / "ofp_flat_short.docx"
    doc.save(path)

    df = load_ofp_data_from_docx(path)
    assert df.empty
