"""Проверка чтения "сырого" Excel-переноса лабораторного отчёта (блок на образец)."""

import openpyxl

from jfunction.raw_excel_io import load_lab_data_from_raw_excel


def _make_raw_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active

    ws["A1"] = "Скважина №300"
    ws["A2"] = "Исходные данные образца №011103002K02H"
    ws["A4"] = "Наименование"
    ws["B4"] = "Значение"
    ws["C4"] = "Водонасыщенность (Sw), д.ед"
    ws["D4"] = "Капиллярное давление, МПа"

    rows = [
        ("Месторождение", "Жанаталап", 1.00, 0.000),
        ("Скважина", 300, 0.95, 0.006),
        ("№ образца", "011103002K02H", 0.88, 0.026),
        ("Глубина, м", 373.94, 0.82, 0.103),
        ("Пористость, %", 31.04, 0.68, 0.233),
        ("Проницаемость по воде, мД", 5.83, 0.40, 1.268),
    ]
    for i, (label, value, sw, pc) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
        ws.cell(row=i, column=3, value=sw)
        ws.cell(row=i, column=4, value=pc)

    ws.cell(row=11, column=1, value="Горизонт")
    ws.cell(row=11, column=2, value="Мел")

    wb.save(path)


def test_reads_block_per_sample_layout(tmp_path):
    path = tmp_path / "raw.xlsx"
    _make_raw_workbook(path)

    df = load_lab_data_from_raw_excel(path)

    assert len(df) == 6
    assert df["well"].iloc[0] == "300"
    assert df["sample"].iloc[0] == "011103002K02H"
    assert df["horizon"].iloc[0] == "Мел"
    assert df["porosity_pct"].iloc[0] == 31.04
    assert df["perm_mD"].iloc[0] == 5.83
    assert list(df["Sw"]) == [1.00, 0.95, 0.88, 0.82, 0.68, 0.40]
    assert list(df["Pc_lab_MPa"]) == [0.000, 0.006, 0.026, 0.103, 0.233, 1.268]
