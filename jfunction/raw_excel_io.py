"""
Чтение "сырых" Excel-файлов лабораторных данных - как переносятся из
Word-отчёта лаборатории (по одному блоку на образец: "Наименование |
Значение | Водонасыщенность (Sw) | Капиллярное давление | ..."), а не
готовой "длинной" таблицы с столбцами well/sample/Sw/Pc_lab_MPa/...

Формат блока (см. examples/, файл J-Function_Gran.xlsx):
    строка "Исходные данные образца №<...>"        - начало блока
    строка "Наименование" / "Значение" / ...       - заголовок таблицы
    далее подряд строки с парами (Sw, Pc) в найденных по заголовку столбцах,
    а в столбцах "Наименование"/"Значение" - метаданные образца
    (Скважина, № образца, Глубина, Пористость, Проницаемость, Горизонт).

Строка "Скважина №<N>" (без "образца") выше блока задаёт номер скважины
для всех блоков до следующей такой строки.
"""

from __future__ import annotations

import numbers
from pathlib import Path

import openpyxl
import pandas as pd

BLOCK_MARKER = "Исходные данные образца"
HEADER_MARKER = "Наименование"
WELL_MARKER = "Скважина №"

LABEL_TARGETS = {
    "№ образца": "sample",
    "Глубина": "depth",
    "Пористость": "porosity_pct",
    "Проницаемость": "perm_mD",
    "Горизонт": "horizon",
}


def _find_block_starts(ws) -> list[int]:
    starts = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and BLOCK_MARKER in v:
            starts.append(r)
    return starts


def _find_well_rows(ws) -> list[tuple[int, str]]:
    rows = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith(WELL_MARKER) and "образца" not in v:
            rows.append((r, v.replace(WELL_MARKER, "").strip()))
    return rows


def _well_for_row(well_rows: list[tuple[int, str]], row: int) -> str | None:
    current = None
    for wr, wv in well_rows:
        if wr <= row:
            current = wv
        else:
            break
    return current


def _parse_sheet(ws) -> list[dict]:
    block_starts = _find_block_starts(ws)
    if not block_starts:
        return []
    block_starts.append(ws.max_row + 1)
    well_rows = _find_well_rows(ws)

    out_rows: list[dict] = []
    for i in range(len(block_starts) - 1):
        bs, be = block_starts[i], block_starts[i + 1]
        well = _well_for_row(well_rows, bs)

        header_row = None
        for r in range(bs, be):
            if ws.cell(row=r, column=1).value == HEADER_MARKER:
                header_row = r
                break
        if header_row is None:
            continue

        col_sw = col_pc = None
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if isinstance(v, str):
                if "одонасыщ" in v:
                    col_sw = c
                if "апил" in v:  # "Капиллярное" и опечатка "Капилярное"
                    col_pc = c
        if col_sw is None or col_pc is None:
            continue

        meta: dict[str, object] = {}
        for r in range(bs, be):
            a = ws.cell(row=r, column=1).value
            if isinstance(a, str):
                for key, name in LABEL_TARGETS.items():
                    if key in a:
                        meta[name] = ws.cell(row=r, column=2).value

        data_pts: list[tuple[float, float]] = []
        r = header_row + 1
        while r < be:
            sw = ws.cell(row=r, column=col_sw).value
            pc = ws.cell(row=r, column=col_pc).value
            if isinstance(sw, numbers.Number) and isinstance(pc, numbers.Number):
                data_pts.append((sw, pc))
                r += 1
            else:
                break

        for sw, pc in data_pts:
            out_rows.append(
                {
                    "well": well,
                    "sample": meta.get("sample"),
                    "horizon": meta.get("horizon"),
                    "depth": meta.get("depth"),
                    "porosity_pct": meta.get("porosity_pct"),
                    "perm_mD": meta.get("perm_mD"),
                    "Sw": sw,
                    "Pc_lab_MPa": pc,
                }
            )

    return out_rows


def load_lab_data_from_raw_excel(path: str | Path) -> pd.DataFrame:
    """
    Извлекает точки измерений из "сырого" Excel-файла лабораторных данных
    (по блоку на образец, как переносится из Word-отчёта). Сканирует все
    листы книги и объединяет найденные блоки.
    """
    wb = openpyxl.load_workbook(str(path), data_only=True)

    rows: list[dict] = []
    for ws in wb.worksheets:
        rows.extend(_parse_sheet(ws))

    if not rows:
        raise ValueError(
            f"В файле {path} не найдено ни одного блока образца "
            f"(строка вида '{BLOCK_MARKER} №...' с таблицей Sw/Pc под ней)."
        )

    return pd.DataFrame(rows)
