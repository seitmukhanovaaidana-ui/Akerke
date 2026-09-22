"""
Петрофизика керна по горизонтам: чтение "Результатов петрофизического
анализа керна" (лист вида "Рез_ан_керна" - № скв., интервал, горизонт,
пористость открытая/полная, насыщенность нефтью/водой, проницаемость по
газу/по воде и т.д.) и построение зависимости k = a*exp(b*Кп) отдельно по
каждому горизонту.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd

COLUMNS = {
    "well": 2,
    "interval": 3,
    "depth": 4,
    "horizon": 5,
    "strat": 6,
    "description": 7,
    "mineral_density": 8,
    "bulk_density": 9,
    "poro_open": 10,
    "poro_full": 11,
    "sat_oil": 12,
    "sat_water": 13,
    "carbonate": 20,
    "perm_gas": 21,
    "perm_water": 22,
}

NUMERIC_COLUMNS = (
    "depth", "mineral_density", "bulk_density", "poro_open", "poro_full",
    "sat_oil", "sat_water", "carbonate", "perm_gas", "perm_water",
)


def _to_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text or text == "-":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def load_core_petro_xlsx(path: str | Path, sheet_name: str | None = None) -> pd.DataFrame:
    """
    Читает лист "Результаты петрофизического анализа керна" (формат
    ПРИЛОЖЕНИЯ 4: №№ скв./интервал/привязанная глубина/горизонт/.../
    пористость открытая-полная/насыщенность нефтью-водой/.../
    проницаемость на газ-на воду).
    """
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    if sheet_name is None:
        for name in wb.sheetnames:
            ws = wb[name]
            if ws.cell(row=1, column=COLUMNS["well"]).value and "скв" in str(
                ws.cell(row=1, column=COLUMNS["well"]).value
            ):
                sheet_name = name
                break
    if sheet_name is None:
        raise ValueError("Не найден лист с результатами петрофизического анализа керна.")

    ws = wb[sheet_name]

    rows: list[dict] = []
    for r in range(4, ws.max_row + 1):
        well = ws.cell(row=r, column=COLUMNS["well"]).value
        if well is None:
            continue
        row = {key: ws.cell(row=r, column=col).value for key, col in COLUMNS.items()}
        for key in NUMERIC_COLUMNS:
            row[key] = _to_float(row[key])
        row["well"] = str(well).strip()
        if row["horizon"] is not None:
            row["horizon"] = str(row["horizon"]).strip()
        rows.append(row)

    return pd.DataFrame(rows)


@dataclass(frozen=True)
class PoroPermFit:
    horizon: str
    n: int
    a: float
    b: float
    r2: float

    def predict(self, poro_pct):
        return self.a * np.exp(self.b * np.asarray(poro_pct, dtype=float))


def group_by_strat(df: pd.DataFrame, strat_col: str = "strat", horizon_col: str = "horizon") -> pd.DataFrame:
    """
    Группирует образцы по столбцу "Стратиграфия" (мел/юра/четверт.),
    который в этом отчёте уже проставлен геологом для каждого образца -
    в отличие от endpoint_cubes.group_by_formation(), которая угадывает
    мел/юра по названию горизонта и может ошибаться на кодах вида
    "K1al2-1" (не содержат кириллического "альб" и т.п., поэтому
    остаются неклассифицированными). Возвращает копию df, где
    horizon_col заменён на strat_col; строки не "мел"/"юра" (например,
    "четверт.") или без стратиграфии отбрасываются.
    """
    if strat_col not in df.columns:
        raise ValueError(f"В данных нет столбца «{strat_col}» (стратиграфия).")
    normalized = df[strat_col].astype(str).str.strip().str.lower()
    out = df.copy()
    out[horizon_col] = normalized
    return out[normalized.isin(["мел", "юра"])]


def _fit_poro_perm(x: np.ndarray, y_perm: np.ndarray) -> tuple[float, float, float]:
    """МНК k=a*exp(b*Кп) по ln(k) от Кп. Возвращает (a, b, r2)."""
    y = np.log(y_perm)
    b, ln_a = np.polyfit(x, y, 1)
    a = float(np.exp(ln_a))
    pred = ln_a + b * x
    ss_res = np.sum((y - pred) ** 2)
    ss_tot = np.sum((y - y.mean()) ** 2)
    r2 = float(1 - ss_res / ss_tot) if ss_tot > 0 else float("nan")
    return a, float(b), r2


def fit_poro_perm_single(
    df: pd.DataFrame,
    label: str,
    poro_col: str = "poro_open",
    perm_col: str = "perm_gas",
    min_samples: int = 3,
) -> dict | None:
    """
    Считает k = a*exp(b*Кп) для ПРОИЗВОЛЬНОГО, заранее отфильтрованного
    подмножества образцов - например, нескольких горизонтов и/или
    стратиграфических единиц, объединённых пользователем вручную (см.
    "Произвольная комбинация горизонтов/стратиграфии" во вкладке
    "Петрофизика по горизонтам"). То же самое, что одна строка
    fit_poro_perm_by_horizon(), но без группировки по столбцу "horizon" -
    вызывающий код сам решает, что попадает в df, а label - это просто
    подпись для отображения (например, "Ю-II + Ю-VI").

    Возвращает None, если после отбраковки пустых/неположительных значений
    осталось меньше min_samples точек.
    """
    sub = df.dropna(subset=[poro_col, perm_col])
    sub = sub[sub[perm_col] > 0]
    if len(sub) < min_samples:
        return None

    x = sub[poro_col].to_numpy()
    a, b, r2 = _fit_poro_perm(x, sub[perm_col].to_numpy())

    return {
        "horizon": label,
        "n": len(sub),
        "a": a,
        "b": b,
        "r2": r2,
        "poro_min": float(sub[poro_col].min()),
        "poro_max": float(sub[poro_col].max()),
        "perm_min": float(sub[perm_col].min()),
        "perm_max": float(sub[perm_col].max()),
    }


def fit_poro_perm_by_horizon(
    df: pd.DataFrame,
    poro_col: str = "poro_open",
    perm_col: str = "perm_gas",
    min_samples: int = 5,
) -> pd.DataFrame:
    """
    Подбирает k = a*exp(b*Кп) отдельно по каждому горизонту (МНК по
    ln(k) от Кп), только для горизонтов, где хватает точек
    (min_samples, по умолчанию 5 - меньше не показательно).

    Возвращает DataFrame: horizon, n, a, b, r2, poro_min, poro_max,
    perm_min, perm_max.
    """
    rows = []
    for horizon, sub in df.groupby("horizon", dropna=True):
        sub = sub.dropna(subset=[poro_col, perm_col])
        sub = sub[sub[perm_col] > 0]
        if len(sub) < min_samples:
            continue

        x = sub[poro_col].to_numpy()
        a, b, r2 = _fit_poro_perm(x, sub[perm_col].to_numpy())

        rows.append(
            {
                "horizon": horizon,
                "n": len(sub),
                "a": a,
                "b": b,
                "r2": r2,
                "poro_min": float(sub[poro_col].min()),
                "poro_max": float(sub[poro_col].max()),
                "perm_min": float(sub[perm_col].min()),
                "perm_max": float(sub[perm_col].max()),
            }
        )

    return pd.DataFrame(rows).sort_values("n", ascending=False).reset_index(drop=True)
