"""
Кубы концевых точек: корреляция Swir/Sor/krwmax (концевые точки Кори) с
пористостью и проницаемостью по каждому горизонту отдельно, и применение
этой корреляции к произвольному массиву значений Кп/k (то есть - к кубу
пористости/проницаемости из геологической модели).

Источник данных - сводная таблица по образцам вида листа "ОФП"
("Таблица 2.4.2 - Относительная проницаемость в системе вода-нефть"):
одна строка на образец/модель, со столбцами скважина, модель, горизонт,
пористость, проницаемость, Swir, Sor, ОФП по воде при Sor (krwmax).

Методика повторяет то, что было сделано вручную в исходном Excel (12
диаграмм рассеяния с линиями тренда на листе "ОФП"): для каждой пары
(горизонт, концевая точка) перебираются 4 типа зависимости от Кп и от k
(линейная, логарифмическая, степенная, экспоненциальная), выбирается та,
что даёт максимальный R².
"""

from __future__ import annotations

import re
from dataclasses import dataclass

import numpy as np
import pandas as pd

SUMMARY_COLUMNS = {
    "well": 2,
    "model": 3,
    "depth": 4,
    "horizon": 6,
    "porosity_pct": 7,
    "perm_mD": 8,
    "Swir": 12,
    "Sor": 13,
    "krwmax": 14,
    "krow_swc": 15,
}

ENDPOINTS = ("Swir", "Sor", "krwmax")
X_VARS = ("porosity_pct", "perm_mD")


def _to_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text or text == "-":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def load_endpoint_summary_xlsx(path: str, sheet_name: str = "ОФП") -> pd.DataFrame:
    """
    Читает сводную таблицу по образцам (лист "ОФП" - Таблица 2.4.2):
    well, model, horizon, porosity_pct, perm_mD, Swir, Sor, krwmax, krow_swc.
    """
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"На листе не найдено «{sheet_name}». Доступные листы: {wb.sheetnames}")
    ws = wb[sheet_name]

    rows: list[dict] = []
    for r in range(5, ws.max_row + 1):
        well = ws.cell(row=r, column=SUMMARY_COLUMNS["well"]).value
        horizon = ws.cell(row=r, column=SUMMARY_COLUMNS["horizon"]).value
        if well is None or horizon is None:
            continue
        row = {key: ws.cell(row=r, column=col).value for key, col in SUMMARY_COLUMNS.items()}
        for key in ("depth", "porosity_pct", "perm_mD", "Swir", "Sor", "krwmax", "krow_swc"):
            row[key] = _to_float(row[key])
        row["well"] = str(well).strip()
        row["horizon"] = str(horizon).strip()
        rows.append(row)

    return pd.DataFrame(rows)


SUMMARY_DOCX_COLUMNS = (
    "well", "model", "horizon", "porosity_pct", "perm_mD", "Swir", "Sor", "krwmax", "krow_swc",
)


def load_endpoint_summary_docx(path: str) -> pd.DataFrame:
    """
    Читает тот же набор образцов (well/model/horizon/Кп/k/Swir/Sor/krwmax),
    что и load_endpoint_summary_xlsx(), но напрямую из "сырого" Word-отчёта
    лаборатории по ОФП - см. ofp_docx_io.load_ofp_data_from_docx() (там же
    описаны оба поддерживаемых формата отчёта). Одна строка на образец/
    модель - кривая Sw/krw/krow сворачивается до её собственных Swir/Sor/
    krwmax/пористости/проницаемости, уже посчитанных при разборе отчёта.

    ВАЖНО: не все форматы отчёта содержат горизонт (лаборатория обычно
    указывает только скважину/модель, а горизонт сопоставляется отдельно
    геологом) - в этом случае столбец horizon будет пустым, и группировка
    по горизонту/"Мел/Юра" не даст результатов для этих образцов, пока
    горизонт не будет проставлен (например, вручную в экспортированной
    таблице или через сводный xlsx, где он уже есть).
    """
    from .ofp_docx_io import load_ofp_data_from_docx

    curve_df = load_ofp_data_from_docx(path)
    if curve_df.empty:
        return pd.DataFrame(columns=list(SUMMARY_DOCX_COLUMNS))

    per_model = curve_df.drop_duplicates(subset=["model"])
    out = pd.DataFrame(
        {
            "well": per_model["well"],
            "model": per_model["model"],
            "horizon": per_model["horizon"] if "horizon" in per_model.columns else "",
            "porosity_pct": per_model["porosity_pct"],
            "perm_mD": per_model["perm_mD"],
            "Swir": per_model["Swir"],
            "Sor": per_model["Sor"],
            "krwmax": per_model["krwmax"],
            "krow_swc": per_model["krow_swc"],
        }
    )
    return out.reset_index(drop=True)


def classify_formation(horizon: str) -> str | None:
    """
    Укрупнённая классификация горизонта до мел/юра по названию (апт,
    альб, неоком, валанжин, "мел..." -> мел; Ю-*, "юра", "...юрский" ->
    юра). Возвращает None, если горизонт не удалось классифицировать
    (пусто, триас и т.п. - триас старше юры, к мелу/юре не относится).
    """
    h = str(horizon).strip().lower()
    if not h:
        return None
    if h.startswith("ю-") or "юр" in h:
        return "юра"
    if any(key in h for key in ("мел", "альб", "апт", "неоком", "валанжин", "готерив", "баррем")):
        return "мел"
    return None


def group_by_formation(df: pd.DataFrame, horizon_col: str = "horizon") -> pd.DataFrame:
    """
    Возвращает копию df, где столбец horizon_col заменён на укрупнённую
    группу "мел"/"юра" (строки, которые не удалось классифицировать,
    отбрасываются).
    """
    out = df.copy()
    out[horizon_col] = out[horizon_col].map(classify_formation)
    return out.dropna(subset=[horizon_col])


FORM_LABELS = {
    "linear": "линейная", "log": "логарифмическая", "power": "степенная",
    "exp": "экспоненциальная", "poly2": "полиномиальная (2-й ст.)",
}


@dataclass(frozen=True)
class Correlation:
    horizon: str
    endpoint: str
    x_var: str
    form: str  # "linear" | "log" | "power" | "exp" | "poly2"
    a: float
    b: float
    r2: float
    n: int
    x_min: float
    x_max: float
    c: float = 0.0  # коэффициент при x² - используется только формой "poly2"

    def predict(self, x):
        x = np.asarray(x, dtype=float)
        if self.form == "linear":
            return self.a + self.b * x
        if self.form == "log":
            return self.a + self.b * np.log(x)
        if self.form == "power":
            return self.a * np.power(x, self.b)
        if self.form == "exp":
            return self.a * np.exp(self.b * x)
        if self.form == "poly2":
            return self.a + self.b * x + self.c * x**2
        raise ValueError(f"Неизвестная форма зависимости: {self.form}")


def _r2(y, pred) -> float:
    ss_res = np.sum((y - pred) ** 2)
    ss_tot = np.sum((y - y.mean()) ** 2)
    return float(1 - ss_res / ss_tot) if ss_tot > 0 else float("nan")


def _fit_form(form: str, x: np.ndarray, y: np.ndarray) -> tuple[float, float, float, float] | None:
    """Возвращает (a, b, c, r2) для заданной формы или None, если форма неприменима к данным."""
    c = 0.0
    try:
        if form == "linear":
            b, a = np.polyfit(x, y, 1)
            pred = a + b * x
        elif form == "log":
            if np.any(x <= 0):
                return None
            b, a = np.polyfit(np.log(x), y, 1)
            pred = a + b * np.log(x)
        elif form == "power":
            if np.any(x <= 0) or np.any(y <= 0):
                return None
            b, ln_a = np.polyfit(np.log(x), np.log(y), 1)
            a = np.exp(ln_a)
            pred = a * np.power(x, b)
        elif form == "exp":
            if np.any(y <= 0):
                return None
            b, ln_a = np.polyfit(x, np.log(y), 1)
            a = np.exp(ln_a)
            pred = a * np.exp(b * x)
        elif form == "poly2":
            if len(x) < 4:  # 3 параметра - на n=3 идеальное совпадение, R² не показателен
                return None
            c2, c1, c0 = np.polyfit(x, y, 2)
            a, b, c = float(c0), float(c1), float(c2)
            pred = a + b * x + c * x**2
        else:
            raise ValueError(form)
    except (np.linalg.LinAlgError, ValueError):
        return None
    return float(a), float(b), float(c), _r2(y, pred)


def fit_all_forms(x, y, horizon: str, endpoint: str, x_var: str) -> dict[str, Correlation]:
    """
    Считает корреляцию ОТДЕЛЬНО для каждой формы (linear/log/power/exp/
    poly2), а не только лучшую по R² - чтобы пользователь мог сам выбрать
    форму тренда для графика (как выбор типа линии тренда в Excel), а не
    полагаться только на автоматический выбор. Формы, неприменимые к
    данным (log/power при x<=0, poly2 при n<4), в словарь не попадают.
    """
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    mask = np.isfinite(x) & np.isfinite(y)
    x, y = x[mask], y[mask]
    if len(x) < 3:
        return {}

    results: dict[str, Correlation] = {}
    for form in ("linear", "log", "power", "exp", "poly2"):
        fit = _fit_form(form, x, y)
        if fit is None:
            continue
        a, b, c, r2 = fit
        results[form] = Correlation(
            horizon=horizon, endpoint=endpoint, x_var=x_var, form=form,
            a=a, b=b, c=c, r2=r2, n=len(x), x_min=float(x.min()), x_max=float(x.max()),
        )
    return results


def fit_best_correlation(x, y, horizon: str, endpoint: str, x_var: str) -> Correlation | None:
    """
    Перебирает linear/log/power/exp (БЕЗ полиномиальной - у неё лишний
    свободный параметр, из-за чего она почти всегда "выигрывает" по R² на
    маленькой выборке чисто за счёт переподгонки, а не реальной связи;
    полином доступен только для ручного выбора формы, см. fit_all_forms())
    и возвращает лучшую по R² (или None, если данных мало).
    """
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    mask = np.isfinite(x) & np.isfinite(y)
    x, y = x[mask], y[mask]
    if len(x) < 3:
        return None

    best = None
    for form in ("linear", "log", "power", "exp"):
        fit = _fit_form(form, x, y)
        if fit is None:
            continue
        a, b, c, r2 = fit
        if best is None or (r2 == r2 and r2 > best[1]):  # r2==r2 отсекает NaN
            best = (form, r2, a, b, c)

    if best is None:
        return None
    form, r2, a, b, c = best
    return Correlation(
        horizon=horizon, endpoint=endpoint, x_var=x_var, form=form,
        a=a, b=b, c=c, r2=r2, n=len(x), x_min=float(x.min()), x_max=float(x.max()),
    )


def fit_endpoint_cubes(df: pd.DataFrame, min_samples: int = 4) -> list[Correlation]:
    """
    Для каждого горизонта и каждой концевой точки (Swir, Sor, krwmax)
    подбирает лучшую корреляцию отдельно от пористости и от проницаемости
    (горизонты с числом образцов меньше min_samples пропускаются).
    """
    results: list[Correlation] = []
    for horizon, sub in df.groupby("horizon"):
        if len(sub) < min_samples:
            continue
        for endpoint in ENDPOINTS:
            for x_var in X_VARS:
                sub2 = sub.dropna(subset=[x_var, endpoint])
                if len(sub2) < min_samples:
                    continue
                corr = fit_best_correlation(sub2[x_var], sub2[endpoint], horizon, endpoint, x_var)
                if corr is not None:
                    results.append(corr)
    return results


def apply_correlation(corr: Correlation, x_values) -> np.ndarray:
    """Применяет корреляцию к массиву значений Кп/k - то есть строит "куб" концевой точки."""
    return corr.predict(x_values)


# --------------------------------------------------------------------------- контроль качества

QUALITY_OK = "ok"
QUALITY_WARNING = "warning"
QUALITY_BAD = "bad"

_BOUNDED_ENDPOINTS = ("Swir", "Sor", "krwmax")


def quality_flags(corr: Correlation) -> list[str]:
    """
    Список замечаний к качеству подбора корреляции:
    - низкий/неопределённый R²;
    - малая выборка (n <= 5), на которой легко получить случайную зависимость;
    - прогноз выходит за физический диапазон [0, 1] уже в пределах диапазона
      обучающих данных (Swir/Sor/krwmax - насыщенности и ОФП, не могут быть
      вне [0, 1]);
    - неустойчивость при небольшой экстраполяции (+-10% от диапазона данных) -
      именно так проявляет себя случай вида Ю-VIб (a~1e14 при отрицательном b).
    """
    flags: list[str] = []

    if corr.r2 != corr.r2:  # NaN
        flags.append("R² не определён")
    elif corr.r2 < 0.5:
        flags.append(f"низкий R² ({corr.r2:.2f})")
    elif corr.r2 < 0.75:
        flags.append(f"средний R² ({corr.r2:.2f})")

    if corr.n <= 5:
        flags.append(f"малая выборка (n={corr.n})")

    if corr.endpoint in _BOUNDED_ENDPOINTS and corr.x_max > corr.x_min:
        xx = np.linspace(corr.x_min, corr.x_max, 20)
        yy = corr.predict(xx)
        if np.any(yy < -0.01) or np.any(yy > 1.01):
            flags.append("прогноз выходит за физический диапазон [0,1] в пределах данных")

        span = corr.x_max - corr.x_min
        x_ext = np.array([corr.x_min - 0.1 * span, corr.x_max + 0.1 * span])
        if corr.form in ("log", "power"):
            x_ext = x_ext[x_ext > 0]
        if len(x_ext):
            y_ext = corr.predict(x_ext)
            if np.any(y_ext < -0.2) or np.any(y_ext > 1.2):
                flags.append("неустойчиво при экстраполяции за пределы диапазона данных")

    return flags


def quality_level(corr: Correlation) -> str:
    """Сводная оценка качества: "bad" (не доверять), "warning" (использовать осторожно), "ok"."""
    flags = quality_flags(corr)
    if any(
        ("физический диапазон" in f) or ("не определён" in f) or ("низкий R²" in f)
        for f in flags
    ):
        return QUALITY_BAD
    if flags:
        return QUALITY_WARNING
    return QUALITY_OK


def check_apply_inputs(corr: Correlation, x_values) -> list[str]:
    """
    Предупреждения при применении корреляции к конкретным значениям Кп/k:
    выход за диапазон обучающих данных (экстраполяция) и нефизичные
    (вне [0,1]) прогнозные значения концевой точки.
    """
    x = np.asarray(x_values, dtype=float)
    warnings: list[str] = []

    out_of_range = (x < corr.x_min) | (x > corr.x_max)
    if np.any(out_of_range):
        warnings.append(
            f"{int(np.sum(out_of_range))} из {len(x)} значений вне диапазона обучающих "
            f"данных [{corr.x_min:.3g}, {corr.x_max:.3g}] - экстраполяция менее надёжна"
        )

    if corr.endpoint in _BOUNDED_ENDPOINTS:
        y = corr.predict(x)
        if np.any(y < 0) or np.any(y > 1):
            warnings.append("часть прогнозных значений выходит за физический диапазон [0, 1]")

    return warnings
